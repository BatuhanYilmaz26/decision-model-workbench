"""Bounded tabular inspection and explicit mapping into unreviewed domain records."""

import csv
import datetime
import hashlib
import io
import math
import zipfile
from xml.etree.ElementTree import ParseError
from pathlib import Path
from .data import dumps, hard_gold, options, validate, validate_questions

MAX_ROWS = 20000
MAX_COLUMNS = 256


def check_upload(content, filename):
    if len(content) > 10 * 1024 * 1024:
        raise ValueError("Upload limit is 10 MiB")
    if Path(filename).suffix.lower() == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if (
                    len(archive.infolist()) > 4096
                    or sum(e.file_size for e in archive.infolist()) > 50 * 1024 * 1024
                ):
                    raise ValueError(
                        "Spreadsheet expanded-size limit is 50 MiB and 4096 archive entries"
                    )
        except zipfile.BadZipFile as exc:
            raise ValueError("Invalid XLSX archive") from exc


def load_xlsx(content):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        if not workbook.sheetnames:
            workbook.close()
            raise ValueError("Workbook needs at least one worksheet")
        return workbook
    except (KeyError, OSError, ParseError, InvalidFileException, IndexError) as exc:
        raise ValueError("Invalid XLSX workbook structure or XML") from exc


def cell(value):
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Nonfinite spreadsheet numbers are not supported")
    return value


def read_table(content, filename, worksheet=None):
    """Return bounded raw table rows. Formula expressions are never evaluated."""
    check_upload(content, filename)
    extension = Path(filename).suffix.lower()
    workbook = None
    try:
        if extension == ".xlsx":
            workbook = load_xlsx(content)
            sheets = workbook.sheetnames
            chosen = worksheet or sheets[0]
            if chosen not in sheets:
                raise ValueError("Selected worksheet does not exist")
            sheet = workbook[chosen]
            if (sheet.max_column or 0) > MAX_COLUMNS or (
                sheet.max_row or 0
            ) > MAX_ROWS + 1:
                raise ValueError("Worksheet exceeds 20,000 rows or 256 columns")
            iterator = sheet.iter_rows(values_only=True)
        elif extension in (".csv", ".tsv"):
            if worksheet:
                raise ValueError("Worksheet selection applies only to XLSX")
            sheets, chosen = [], None
            iterator = csv.reader(
                io.StringIO(content.decode("utf-8-sig"), newline=""),
                delimiter="\t" if extension == ".tsv" else ",",
                strict=True,
            )
        else:
            raise ValueError("Column mapping supports CSV, TSV and XLSX")
        raw_headers = next(iterator, None)
        if not raw_headers:
            raise ValueError("Spreadsheet is empty")
        if len(raw_headers) > MAX_COLUMNS or any(
            v is None or not str(v).strip() for v in raw_headers
        ):
            raise ValueError(
                "Table headers must be nonempty; at most 256 columns are supported"
            )
        headers = [str(v).strip() for v in raw_headers]
        if len(set(headers)) != len(headers):
            raise ValueError("Table headers must be unique")
        rows = []
        for number, values in enumerate(iterator, 2):
            if number > MAX_ROWS + 1:
                raise ValueError("Table exceeds 20,000 source rows")
            if not values or all(v is None or v == "" for v in values):
                continue
            if len(values) != len(headers):
                raise ValueError(
                    f"Source row {number} has {len(values)} cells; expected {len(headers)}"
                )
            rows.append((number, {k: cell(v) for k, v in zip(headers, values)}))
        if not rows:
            raise ValueError("No records found")
        return {
            "worksheets": sheets,
            "worksheet": chosen,
            "columns": headers,
            "rows": rows,
        }
    except (UnicodeError, csv.Error) as exc:
        raise ValueError(
            "Read failed; use UTF-8 CSV/TSV with consistently quoted columns"
        ) from exc
    except (zipfile.BadZipFile, KeyError, ParseError) as exc:
        raise ValueError("Invalid spreadsheet content") from exc
    finally:
        if workbook is not None:
            workbook.close()


def inspect_table(content, filename, worksheet=None):
    if Path(filename).suffix.lower() == ".xlsx" and worksheet is None:
        check_upload(content, filename)
        workbook = load_xlsx(content)
        try:
            sheets = workbook.sheetnames
        finally:
            workbook.close()
        try:
            table = read_table(content, filename)
        except ValueError as exc:
            return {
                "worksheets": sheets,
                "worksheet": sheets[0],
                "columns": [],
                "row_count": 0,
                "preview": [],
                "column_summary": {},
                "worksheet_error": str(exc),
                "upload_sha256": hashlib.sha256(content).hexdigest(),
            }
    else:
        table = read_table(content, filename, worksheet)
    categories = {}
    for column in table["columns"]:
        counts, missing = {}, 0
        for _, row in table["rows"]:
            value = row[column]
            if value is None or value == "":
                missing += 1
            else:
                key = str(value)
                if key in counts or len(counts) < 100:
                    counts[key] = counts.get(key, 0) + 1
        categories[column] = {
            "missing": missing,
            "sample_values": [v for v in list(counts)[:100] if len(v) <= 200],
            "values_are_sampled": len(counts) >= 100
            or any(len(v) > 200 for v in counts),
        }
    return {k: v for k, v in table.items() if k != "rows"} | {
        "row_count": len(table["rows"]),
        "preview": [
            {k: v[:500] if isinstance(v, str) else v for k, v in row.items()}
            for _, row in table["rows"][:10]
        ],
        "preview_string_limit": 500,
        "column_summary": categories,
        "upload_sha256": hashlib.sha256(content).hexdigest(),
    }


def map_table(content, filename, questions, mapping):
    validate_questions(questions)
    allowed = {
        "worksheet",
        "state_columns",
        "state_mode",
        "id_column",
        "group_column",
        "labels",
        "value_maps",
        "origin",
        "temporal_columns",
    }
    if not isinstance(mapping, dict) or set(mapping) - allowed:
        raise ValueError("Unexpected column-mapping fields")
    origin = mapping.get("origin", "real")
    if origin not in ("real", "synthetic"):
        raise ValueError("Source origin must be real or synthetic")
    table = read_table(content, filename, mapping.get("worksheet"))
    columns = set(table["columns"])
    from .temporal import FIELDS

    temporal_columns = mapping.get("temporal_columns", {})
    if not isinstance(temporal_columns, dict) or (
        temporal_columns
        and (
            set(temporal_columns) != set(FIELDS)
            or any(
                not isinstance(v, str) or v not in columns
                for v in temporal_columns.values()
            )
        )
    ):
        raise ValueError(
            "Map all three temporal fields to existing timestamp columns, or omit temporal_columns"
        )
    state_columns = mapping.get("state_columns")
    if (
        not isinstance(state_columns, list)
        or not 1 <= len(state_columns) <= 64
        or len(set(state_columns)) != len(state_columns)
        or not set(state_columns) <= columns
    ):
        raise ValueError("Select 1–64 unique input columns from the table")
    mode = mapping.get("state_mode", "text")
    if mode not in ("text", "json") or (mode == "text" and len(state_columns) != 1):
        raise ValueError("Use text for one input column or JSON for multiple columns")
    for name in ("id_column", "group_column"):
        if mapping.get(name) is not None and mapping[name] not in columns:
            raise ValueError(name + " must identify an existing column")
    labels, aliases = mapping.get("labels", {}), mapping.get("value_maps", {})
    if (
        not isinstance(labels, dict)
        or not set(labels) <= set(questions)
        or any(v is not None and v not in columns for v in labels.values())
    ):
        raise ValueError(
            "Map project questions to existing label columns or leave them unassigned"
        )
    if not isinstance(aliases, dict) or not set(aliases) <= set(questions):
        raise ValueError("Value mappings must use project question identifiers")
    for name, values in aliases.items():
        if not isinstance(values, dict) or any(
            not isinstance(k, str)
            or not isinstance(v, str)
            or v not in options(questions[name])
            for k, v in values.items()
        ):
            raise ValueError(
                "Label value mappings must point to allowed project outcomes"
            )
    upload_hash = hashlib.sha256(content).hexdigest()
    sheet_hash = hashlib.sha256((table["worksheet"] or "").encode()).hexdigest()[:8]
    result = []
    for number, source in table["rows"]:
        generated_id = f"import-{upload_hash[:16]}-{sheet_hash}-{number}"

        def identifier(field, fallback):
            value = source[mapping[field]] if mapping.get(field) else fallback
            if value is None or not str(value).strip():
                raise ValueError(f"Source row {number}: {field} is blank")
            return str(value).strip()

        record_id = identifier("id_column", generated_id)
        state = (
            source[state_columns[0]]
            if mode == "text"
            else {k: source[k] for k in state_columns}
        )
        if mode == "text" and state is not None:
            state = str(state)
        if (mode == "text" and not state) or (
            mode == "json" and all(v is None or v == "" for v in state.values())
        ):
            raise ValueError(f"Source row {number}: input evidence is blank")
        if len(dumps(state).encode()) > 64000:
            raise ValueError(f"Source row {number}: input evidence exceeds 64 KB")
        targets = {}
        for name, column in labels.items():
            value = source[column] if column else None
            if value is None or value == "":
                continue
            targets[name] = aliases.get(name, {}).get(str(value), value)
        try:
            gold = hard_gold({k: questions[k] for k in targets}, targets)
        except ValueError as exc:
            raise ValueError(f"Source row {number}: {exc}") from exc
        row = {
            "id": record_id,
            "group_id": identifier("group_column", record_id),
            "state": state,
            "questions": questions,
            "gold": gold,
            "metadata": {
                "origin": origin,
                "reviewed": False,
                "source": Path(filename).name,
                "source_row": number,
                "worksheet": table["worksheet"],
            },
        }
        if temporal_columns:
            row["metadata"]["temporal"] = {
                field: source[column] for field, column in temporal_columns.items()
            }
        result.append(validate(row, allow_missing_labels=True))
    if len({r["id"] for r in result}) != len(result):
        raise ValueError("Duplicate record identifiers; choose a unique ID column")
    return result
